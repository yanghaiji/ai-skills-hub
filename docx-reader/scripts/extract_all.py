"Extract both text and tables from Word (.docx) files in one pass."

import argparse
import os
import sys


def extract_all(docx_path: str, output: str | None = None) -> None:
    try:
        from docx import Document
    except ImportError:
        print("Error: python-docx is not installed. Run: pip install python-docx", file=sys.stderr)
        sys.exit(1)

    try:
        doc = Document(docx_path)
    except Exception as e:
        print(f"Error reading docx: {e}", file=sys.stderr)
        sys.exit(1)

    file_name = os.path.basename(docx_path)
    lines: list[str] = []

    def out(text: str = ""):
        lines.append(text)

    out(f"# Word: {file_name}")
    out()

    heading_map = {
        "Heading 1": "# ",
        "Heading 2": "## ",
        "Heading 3": "### ",
        "Heading 4": "#### ",
        "Title": "# ",
    }

    body_children = list(doc.element.body)
    element_table_map: dict[int, int] = {}

    for bi, bc in enumerate(body_children):
        for ti, table in enumerate(doc.tables):
            if bc is table._element:
                element_table_map[bi] = ti
                break

    table_count = 0

    for bi, bc in enumerate(body_children):
        if bi in element_table_map:
            ti = element_table_map[bi]
            table = doc.tables[ti]
            table_count += 1
            out(f"\n--- Table {table_count} ---\n")
            print_markdown_table(table, out)
            out()
        else:
            tag = bc.tag.split("}")[-1] if "}" in bc.tag else bc.tag
            if tag == "p":
                para = None
                for p in doc.paragraphs:
                    if p._element is bc:
                        para = p
                        break
                if para is not None:
                    style_name = para.style.name if para.style else ""
                    prefix = heading_map.get(style_name, "")
                    text = para.text.strip()
                    if text:
                        out(f"{prefix}{text}")
                        out()

    if table_count > 0:
        out(f"\nTotal: {table_count} table(s) extracted.")

    result = "\n".join(lines)

    if output:
        with open(output, "w", encoding="utf-8") as f:
            f.write(result)
        print(f"Output saved to: {output}")
    else:
        print(result)


def print_markdown_table(table, out=print) -> None:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append(cells)

    if not rows:
        out("[Empty table]")
        return

    col_count = max(len(row) for row in rows)
    for row in rows:
        while len(row) < col_count:
            row.append("")

    col_widths = [0] * col_count
    for row in rows:
        for i, cell in enumerate(row):
            col_widths[i] = max(col_widths[i], len(cell))
    col_widths = [max(w, 3) for w in col_widths]

    header = rows[0]
    header_line = "| " + " | ".join(cell.ljust(col_widths[i]) for i, cell in enumerate(header)) + " |"
    separator = "| " + " | ".join("-" * col_widths[i] for i in range(col_count)) + " |"
    out(header_line)
    out(separator)

    for row in rows[1:]:
        line = "| " + " | ".join(cell.ljust(col_widths[i]) for i, cell in enumerate(row)) + " |"
        out(line)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Extract text and tables from Word (.docx) files")
    parser.add_argument("docx_path", help="Path to the .docx file")
    parser.add_argument("--output", "-o", default=None, help="Save output to file instead of stdout")
    args = parser.parse_args()

    extract_all(args.docx_path, args.output)
